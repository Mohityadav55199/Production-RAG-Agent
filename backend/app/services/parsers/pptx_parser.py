import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.services.parsers.base import DocumentParser, ExtractedPage, ExtractionResult
from app.services.parsers.figures import RawImage, build_figures, caption_figures, merge_captions
from app.utils.logging import get_logger

logger = get_logger(__name__)

MAX_IMAGE_CANDIDATES = 150


def _iter_shapes(shapes):
    """Every shape on a slide, including the contents of groups."""
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _slide_images(slide) -> list[bytes]:
    blobs = []
    for shape in _iter_shapes(slide.shapes):
        try:
            if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.PLACEHOLDER) and hasattr(shape, "image"):
                blobs.append(shape.image.blob)
        except Exception as exc:  # placeholders without a picture raise; that is normal
            logger.debug("No image in shape: %s", exc)
    return blobs


class PptxParser(DocumentParser):
    def parse(self, content: bytes, filename: str) -> ExtractionResult:
        try:
            prs = Presentation(io.BytesIO(content))
        except Exception as exc:
            raise ValueError(f"Cannot open PPTX '{filename}': {exc}") from exc

        pages: list[ExtractedPage] = []
        raw_images: list[RawImage] = []
        cursor = 0
        total_pages = len(prs.slides)

        for i, slide in enumerate(prs.slides):
            if len(raw_images) < MAX_IMAGE_CANDIDATES:
                raw_images += [RawImage(i + 1, blob) for blob in _slide_images(slide)]

            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    slide_text.append(shape.text.strip())

            page_text = "\n".join(slide_text).strip()
            if not page_text:
                continue

            start = cursor
            end = start + len(page_text)

            pages.append(ExtractedPage(
                page_number=i + 1,
                text=page_text,
                char_start=start,
                char_end=end
            ))
            cursor = end + 2  # account for \n\n

        images = build_figures(raw_images)
        if images:
            caption_figures(images, hint=f"An image from a slide in the presentation '{filename}'")
            pages = merge_captions(pages, images)

        if not pages:
            raise ValueError(f"PPTX '{filename}' contains no extractable text.")

        full_text = "\n\n".join(p.text for p in pages)

        return ExtractionResult(
            pages=pages,
            full_text=full_text,
            page_count=total_pages,
            char_count=len(full_text),
            file_type="pptx",
            parser_used="python-pptx",
            ocr_used=False,
            ocr_engine=None,
            extraction_method="native",
            images=images,
        )
