"""Pictures found in documents: extraction, filtering, captions, and image files as documents."""

import io

import pytest
from docx import Document as DocxDocument
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.config import get_settings
from app.services.parsers import figures
from app.services.parsers.docx_parser import DocxParser
from app.services.parsers.factory import get_parser_for_file
from app.services.parsers.figures import RawImage, build_figures, merge_captions
from app.services.parsers.image_parser import ImageParser
from app.services.parsers.pdf_parser import PDFParser
from app.services.parsers.pptx_parser import PptxParser
from tests.images import SENTENCE, encoded, pdf_with_images, png


@pytest.fixture
def captions(monkeypatch):
    """Replace Gemini captioning with a recorder that returns a predictable caption."""
    calls = []

    def fake(jpeg, hint=""):
        calls.append(hint)
        return f"Caption number {len(calls)}"

    monkeypatch.setattr(figures, "describe_image_sync", fake)
    return calls


@pytest.fixture
def no_captions(monkeypatch):
    monkeypatch.setattr(figures, "describe_image_sync", lambda jpeg, hint="": None)


class TestBuildFigures:
    def test_icons_banners_and_junk_are_dropped(self):
        raw = [
            RawImage(1, png(300, 300)),                 # kept
            RawImage(1, png(40, 40)),                   # icon
            RawImage(1, png(1400, 60)),                 # thin banner / rule
            RawImage(1, b"definitely not an image"),    # junk
            RawImage(2, png(100, 100)),                 # below the minimum side
        ]
        kept = build_figures(raw)
        assert len(kept) == 1 and kept[0].page_number == 1

    def test_duplicates_are_kept_once(self):
        same = png(300, 300)
        assert len(build_figures([RawImage(1, same), RawImage(2, same), RawImage(3, same)])) == 1

    def test_the_largest_are_kept_when_over_the_cap_and_returned_in_page_order(self):
        raw = [RawImage(p, png(200 + 40 * p, 200 + 40 * p)) for p in range(1, 7)]
        kept = build_figures(raw, max_count=3)
        assert [f.page_number for f in kept] == [4, 5, 6]            # the three biggest, in page order

    def test_order_within_a_page_is_recorded(self):
        kept = build_figures([RawImage(1, png(300, 300)), RawImage(1, png(310, 310)), RawImage(2, png(320, 320))])
        assert [(f.page_number, f.position) for f in kept] == [(1, 0), (1, 1), (2, 0)]

    def test_pages_can_be_skipped(self):
        assert build_figures([RawImage(1, png(300, 300))], skip_pages={1}) == []


class TestMergeCaptions:
    def test_captions_are_appended_to_their_page_and_offsets_stay_consistent(self, captions):
        from app.services.parsers.base import ExtractedPage

        pages = [ExtractedPage(1, "First page text", 0, 15), ExtractedPage(2, "Second", 17, 23)]
        images = build_figures([RawImage(2, png(300, 300))])
        figures.caption_figures(images)
        merged = merge_captions(pages, images)
        assert "[Figure on page 2: Caption number 1]" in merged[1].text and merged[0].text == "First page text"
        full = "\n\n".join(p.text for p in merged)
        assert all(full[p.char_start:p.char_end] == p.text for p in merged)

    def test_a_picture_only_page_becomes_a_page(self, captions):
        from app.services.parsers.base import ExtractedPage

        images = build_figures([RawImage(3, png(300, 300))])
        figures.caption_figures(images)
        merged = merge_captions([ExtractedPage(1, "Only text", 0, 9)], images)
        assert [p.page_number for p in merged] == [1, 3] and "Caption number 1" in merged[1].text

    def test_nothing_changes_without_captions(self, no_captions):
        from app.services.parsers.base import ExtractedPage

        pages = [ExtractedPage(1, "text", 0, 4)]
        images = build_figures([RawImage(1, png(300, 300))])
        figures.caption_figures(images)
        assert merge_captions(pages, images) is pages

    def test_captioning_is_limited_per_document(self, captions, monkeypatch):
        monkeypatch.setattr(get_settings(), "IMAGE_CAPTION_MAX_PER_DOC", 2)
        images = build_figures([RawImage(p, png(250 + p, 250 + p)) for p in range(1, 6)])
        figures.caption_figures(images)
        assert len(captions) == 2 and sum(1 for i in images if i.caption) == 2


class TestPdf:
    def test_a_real_figure_is_extracted_and_a_tiny_icon_is_not(self, no_captions):
        pdf = pdf_with_images([{"text": SENTENCE * 3, "images": [png(400, 300), png(40, 40)]}])
        result = PDFParser().parse(pdf, "report.pdf")
        assert len(result.images) == 1
        image = result.images[0]
        assert image.page_number == 1 and image.kind == "figure"
        assert image.display[:2] == b"\xff\xd8" and image.thumb[:2] == b"\xff\xd8"
        with Image.open(io.BytesIO(image.thumb)) as thumb:
            assert max(thumb.size) <= 320

    def test_images_are_found_on_the_right_pages(self, no_captions):
        pdf = pdf_with_images([
            {"text": SENTENCE * 3, "images": []},
            {"text": SENTENCE * 3, "images": [png(380, 280)]},
            {"text": SENTENCE * 3, "images": [png(390, 290)]},
        ])
        assert [i.page_number for i in PDFParser().parse(pdf, "r.pdf").images] == [2, 3]

    def test_a_pdf_without_pictures_has_none(self, no_captions):
        assert PDFParser().parse(pdf_with_images([{"text": SENTENCE * 3}]), "plain.pdf").images == []

    def test_captions_become_searchable_text_on_that_page(self, captions):
        pdf = pdf_with_images([{"text": SENTENCE * 3, "images": [png(400, 300)]}])
        result = PDFParser().parse(pdf, "report.pdf")
        assert "[Figure on page 1: Caption number 1]" in result.full_text
        assert result.images[0].caption == "Caption number 1"
        assert all(result.full_text[p.char_start:p.char_end] == p.text for p in result.pages)

    def test_a_slide_style_page_with_only_a_picture_is_still_indexed_through_its_caption(self, captions):
        pdf = pdf_with_images([{"text": SENTENCE * 3}, {"text": "", "images": [png(400, 300)]}])
        result = PDFParser().parse(pdf, "deck.pdf")
        assert [p.page_number for p in result.pages] == [1, 2] and "Caption number 1" in result.pages[1].text

    def test_scanned_pages_are_not_treated_as_figures(self, no_captions):
        """A scan is one big photo of the page. It is read by OCR and must not be stored as a figure."""
        parser = PDFParser()
        parser.ocr_service = type("OCR", (), {"extract_text": lambda self, img: ("Text read from the scan", "fake")})()
        result = parser.parse(pdf_with_images([{"text": "", "images": [png(500, 600)]}]), "scan.pdf")
        assert result.ocr_used and result.images == []

    def test_a_pdf_with_nothing_readable_is_still_an_error(self, no_captions):
        parser = PDFParser()
        parser.ocr_service = type("OCR", (), {"extract_text": lambda self, img: ("", "fake")})()
        with pytest.raises(ValueError, match="no extractable text"):
            parser.parse(pdf_with_images([{"text": "", "images": []}]), "empty.pdf")


class TestPowerPointAndWord:
    @staticmethod
    def deck(slides):
        prs = Presentation()
        for text, pics in slides:
            slide = prs.slides.add_slide(prs.slide_layouts[5])        # title only
            slide.shapes.title.text = text
            for data in pics:
                slide.shapes.add_picture(io.BytesIO(data), Inches(1), Inches(2), width=Inches(4))
        out = io.BytesIO()
        prs.save(out)
        return out.getvalue()

    def test_slide_pictures_are_extracted_with_their_slide_number(self, no_captions):
        result = PptxParser().parse(self.deck([("Intro", []), ("Results", [png(400, 300)])]), "deck.pptx")
        assert [(i.page_number, i.kind) for i in result.images] == [(2, "figure")]

    def test_a_slide_with_only_a_picture_is_indexed_through_its_caption(self, captions):
        result = PptxParser().parse(self.deck([("", [png(400, 300)]), ("Next", [])]), "deck.pptx")
        assert "Caption number 1" in result.full_text

    def test_word_pictures_are_extracted(self, no_captions):
        doc = DocxDocument()
        doc.add_paragraph("A report with a chart. " * 5)
        doc.add_picture(io.BytesIO(png(400, 300)))
        out = io.BytesIO()
        doc.save(out)
        result = DocxParser().parse(out.getvalue(), "report.docx")
        assert len(result.images) == 1 and result.images[0].page_number == 1

    def test_word_captions_are_added_to_the_text(self, captions):
        doc = DocxDocument()
        doc.add_paragraph("A report with a chart. " * 5)
        doc.add_picture(io.BytesIO(png(400, 300)))
        out = io.BytesIO()
        doc.save(out)
        result = DocxParser().parse(out.getvalue(), "report.docx")
        assert "[Figure on page 1: Caption number 1]" in result.full_text


class TestImageFiles:
    @pytest.fixture
    def parser(self, monkeypatch):
        parser = ImageParser()
        state = {"ocr": "Total: 42 units", "caption": "A bar chart of unit sales"}
        parser.ocr_service = type("OCR", (), {"extract_text": lambda self, img: (state["ocr"], "fake-ocr")})()
        monkeypatch.setattr("app.services.parsers.image_parser.describe_image_sync", lambda jpeg, hint="": state["caption"])
        parser.state = state
        return parser

    def test_description_and_ocr_both_become_searchable_text(self, parser):
        result = parser.parse(png(400, 300), "sales.png")
        assert "Description: A bar chart of unit sales" in result.full_text
        assert "Total: 42 units" in result.full_text and "sales.png" in result.full_text
        assert result.extraction_method == "ocr+vision" and result.ocr_used

    def test_the_picture_itself_is_kept_for_thumbnails(self, parser):
        result = parser.parse(png(400, 300), "sales.png")
        assert len(result.images) == 1 and result.images[0].kind == "image"
        assert result.images[0].caption == "A bar chart of unit sales"

    def test_a_photo_with_no_text_works_when_it_can_be_described(self, parser):
        parser.state["ocr"] = ""
        result = parser.parse(png(400, 300), "photo.jpg")
        assert "Description:" in result.full_text and "OCR" not in result.full_text
        assert result.extraction_method == "vision" and not result.ocr_used

    def test_text_only_images_still_work_without_a_description(self, parser):
        parser.state["caption"] = None
        result = parser.parse(png(400, 300), "scan.png")
        assert "Description:" not in result.full_text and result.extraction_method == "ocr"

    def test_nothing_readable_and_no_description_is_a_clear_error(self, parser):
        parser.state["ocr"], parser.state["caption"] = "", None
        with pytest.raises(ValueError, match="GOOGLE_API_KEY"):
            parser.parse(png(400, 300), "blank.png")

    def test_a_failing_ocr_does_not_matter_when_a_description_exists(self, parser):
        def boom(self, img):
            raise RuntimeError("tesseract missing")

        parser.ocr_service = type("OCR", (), {"extract_text": boom})()
        assert "Description:" in parser.parse(png(400, 300), "x.png").full_text

    def test_a_non_image_is_rejected(self, parser):
        with pytest.raises(ValueError, match="Cannot open Image"):
            parser.parse(b"not an image", "fake.png")

    def test_image_extensions_use_the_image_parser(self):
        for name in ("a.png", "b.jpg", "c.jpeg", "d.webp", "e.bmp", "f.tiff"):
            assert isinstance(get_parser_for_file(name), ImageParser)
